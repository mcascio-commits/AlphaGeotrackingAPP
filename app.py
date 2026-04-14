from flask import Flask, render_template, redirect, url_for, request, flash, jsonify, send_file
from flask_sqlalchemy import SQLAlchemy
from flask_login import LoginManager, UserMixin, login_user, logout_user, login_required, current_user
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from datetime import datetime
from functools import wraps
import os, io, csv, zipfile

# ── Config ──────────────────────────────────────────────────────────────────
app = Flask(__name__)
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY', 'change-this-in-production-xyz987')
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///autotrack.db'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['UPLOAD_FOLDER'] = os.path.join(os.path.dirname(__file__), 'static', 'uploads')
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024  # 16 MB max photo
ALLOWED_EXTENSIONS = {'png', 'jpg', 'jpeg', 'gif', 'webp'}

db = SQLAlchemy(app)
login_manager = LoginManager(app)
login_manager.login_view = 'login'
login_manager.login_message = 'Veuillez vous connecter pour accéder à cette page.'

os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

# ── Models ───────────────────────────────────────────────────────────────────
class User(UserMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(80), unique=True, nullable=False)
    password_hash = db.Column(db.String(200), nullable=False)
    is_admin = db.Column(db.Boolean, default=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    fiches = db.relationship('Fiche', backref='auteur', lazy=True)

    def set_password(self, pw): self.password_hash = generate_password_hash(pw)
    def check_password(self, pw): return check_password_hash(self.password_hash, pw)

class Partenaire(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    nom = db.Column(db.String(100), unique=True, nullable=False)
    actif = db.Column(db.Boolean, default=True)
    fiches = db.relationship('Fiche', backref='partenaire', lazy=True)

class TypeClient(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    nom = db.Column(db.String(100), unique=True, nullable=False)
    actif = db.Column(db.Boolean, default=True)
    interventions = db.relationship('TypeIntervention', backref='type_client', lazy=True)
    fiches = db.relationship('Fiche', backref='type_client', lazy=True)

class TypeVehicule(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    nom = db.Column(db.String(100), unique=True, nullable=False)
    actif = db.Column(db.Boolean, default=True)
    fiches = db.relationship('Fiche', backref='type_vehicule', lazy=True)

class TypeIntervention(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    nom = db.Column(db.String(100), nullable=False)
    type_client_id = db.Column(db.Integer, db.ForeignKey('type_client.id'), nullable=False)
    actif = db.Column(db.Boolean, default=True)
    fiches = db.relationship('Fiche', backref='type_intervention', lazy=True)

class Fiche(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    nom_client = db.Column(db.String(150), nullable=False)
    plaque = db.Column(db.String(30), nullable=False)
    commentaire = db.Column(db.Text, nullable=True)
    photo_filename = db.Column(db.String(200), nullable=True)
    date_encodage = db.Column(db.DateTime, default=datetime.utcnow)
    partenaire_id = db.Column(db.Integer, db.ForeignKey('partenaire.id'), nullable=False)
    type_client_id = db.Column(db.Integer, db.ForeignKey('type_client.id'), nullable=False)
    type_vehicule_id = db.Column(db.Integer, db.ForeignKey('type_vehicule.id'), nullable=False)
    type_intervention_id = db.Column(db.Integer, db.ForeignKey('type_intervention.id'), nullable=False)
    user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)

# ── Helpers ──────────────────────────────────────────────────────────────────
@login_manager.user_loader
def load_user(user_id): return User.query.get(int(user_id))

def admin_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not current_user.is_authenticated or not current_user.is_admin:
            flash('Accès réservé aux administrateurs.', 'danger')
            return redirect(url_for('dashboard'))
        return f(*args, **kwargs)
    return decorated

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

# ── Auth ─────────────────────────────────────────────────────────────────────
@app.route('/', methods=['GET', 'POST'])
def login():
    if current_user.is_authenticated:
        return redirect(url_for('dashboard'))
    if request.method == 'POST':
        user = User.query.filter_by(username=request.form['username']).first()
        if user and user.check_password(request.form['password']):
            login_user(user, remember=True)
            return redirect(url_for('dashboard'))
        flash('Identifiant ou mot de passe incorrect.', 'danger')
    return render_template('login.html')

@app.route('/logout')
@login_required
def logout():
    logout_user()
    return redirect(url_for('login'))

# ── Dashboard ────────────────────────────────────────────────────────────────
@app.route('/dashboard')
@login_required
def dashboard():
    fiches = Fiche.query.order_by(Fiche.date_encodage.desc()).limit(20).all()
    total = Fiche.query.count()
    today = Fiche.query.filter(db.func.date(Fiche.date_encodage) == datetime.utcnow().date()).count()
    return render_template('dashboard.html', fiches=fiches, total=total, today=today)

# ── Nouvelle fiche ────────────────────────────────────────────────────────────
@app.route('/fiche/nouvelle', methods=['GET', 'POST'])
@login_required
def nouvelle_fiche():
    partenaires = Partenaire.query.filter_by(actif=True).all()
    types_client = TypeClient.query.filter_by(actif=True).all()
    types_vehicule = TypeVehicule.query.filter_by(actif=True).all()

    if request.method == 'POST':
        photo_filename = None
        if 'photo' in request.files:
            file = request.files['photo']
            if file and file.filename and allowed_file(file.filename):
                filename = secure_filename(f"{datetime.utcnow().strftime('%Y%m%d%H%M%S')}_{file.filename}")
                file.save(os.path.join(app.config['UPLOAD_FOLDER'], filename))
                photo_filename = filename

        fiche = Fiche(
            nom_client=request.form['nom_client'],
            plaque=request.form['plaque'].upper(),
            commentaire=request.form.get('commentaire', ''),
            photo_filename=photo_filename,
            partenaire_id=request.form['partenaire_id'],
            type_client_id=request.form['type_client_id'],
            type_vehicule_id=request.form['type_vehicule_id'],
            type_intervention_id=request.form['type_intervention_id'],
            user_id=current_user.id
        )
        db.session.add(fiche)
        db.session.commit()
        flash('Fiche enregistrée avec succès !', 'success')
        return redirect(url_for('dashboard'))

    return render_template('nouvelle_fiche.html', partenaires=partenaires,
                           types_client=types_client, types_vehicule=types_vehicule)

@app.route('/api/interventions/<int:type_client_id>')
@login_required
def get_interventions(type_client_id):
    interventions = TypeIntervention.query.filter_by(type_client_id=type_client_id, actif=True).all()
    return jsonify([{'id': i.id, 'nom': i.nom} for i in interventions])

# ── Détail fiche ──────────────────────────────────────────────────────────────
@app.route('/fiche/<int:fiche_id>')
@login_required
def detail_fiche(fiche_id):
    fiche = Fiche.query.get_or_404(fiche_id)
    return render_template('detail_fiche.html', fiche=fiche)

@app.route('/fiche/<int:fiche_id>/supprimer', methods=['POST'])
@login_required
@admin_required
def supprimer_fiche(fiche_id):
    fiche = Fiche.query.get_or_404(fiche_id)
    if fiche.photo_filename:
        path = os.path.join(app.config['UPLOAD_FOLDER'], fiche.photo_filename)
        if os.path.exists(path): os.remove(path)
    db.session.delete(fiche)
    db.session.commit()
    flash('Fiche supprimée.', 'success')
    return redirect(url_for('dashboard'))

# ── Export ────────────────────────────────────────────────────────────────────
@app.route('/export', methods=['GET', 'POST'])
@login_required
def export():
    partenaires = Partenaire.query.filter_by(actif=True).all()
    types_client = TypeClient.query.filter_by(actif=True).all()

    if request.method == 'POST':
        q = Fiche.query
        if request.form.get('partenaire_id'):
            q = q.filter_by(partenaire_id=request.form['partenaire_id'])
        if request.form.get('type_client_id'):
            q = q.filter_by(type_client_id=request.form['type_client_id'])
        if request.form.get('date_debut'):
            q = q.filter(Fiche.date_encodage >= datetime.strptime(request.form['date_debut'], '%Y-%m-%d'))
        if request.form.get('date_fin'):
            from datetime import timedelta
            d = datetime.strptime(request.form['date_fin'], '%Y-%m-%d') + timedelta(days=1)
            q = q.filter(Fiche.date_encodage < d)
        fiches = q.order_by(Fiche.date_encodage.desc()).all()

        fmt = request.form.get('format', 'csv')
        if fmt == 'csv':
            output = io.StringIO()
            writer = csv.writer(output, delimiter=';')
            writer.writerow(['ID','Date','Heure','Partenaire','Type Client','Nom Client',
                             'Plaque','Type Véhicule','Type Intervention','Commentaire','Encodé par'])
            for f in fiches:
                writer.writerow([f.id, f.date_encodage.strftime('%d/%m/%Y'),
                                 f.date_encodage.strftime('%H:%M'),
                                 f.partenaire.nom, f.type_client.nom, f.nom_client,
                                 f.plaque, f.type_vehicule.nom, f.type_intervention.nom,
                                 f.commentaire or '', f.auteur.username])
            output.seek(0)
            return send_file(io.BytesIO(output.getvalue().encode('utf-8-sig')),
                             mimetype='text/csv',
                             as_attachment=True,
                             download_name=f'autotrack_{datetime.now().strftime("%Y%m%d_%H%M")}.csv')
        else:  # xlsx
            try:
                from openpyxl import Workbook
                from openpyxl.styles import Font, PatternFill, Alignment
                wb = Workbook()
                ws = wb.active
                ws.title = 'Fiches'
                headers = ['ID','Date','Heure','Partenaire','Type Client','Nom Client',
                           'Plaque','Type Véhicule','Type Intervention','Commentaire','Encodé par']
                for col, h in enumerate(headers, 1):
                    cell = ws.cell(row=1, column=col, value=h)
                    cell.font = Font(bold=True, color='FFFFFF')
                    cell.fill = PatternFill('solid', fgColor='1a1a2e')
                    cell.alignment = Alignment(horizontal='center')
                for row, f in enumerate(fiches, 2):
                    ws.append([f.id, f.date_encodage.strftime('%d/%m/%Y'),
                               f.date_encodage.strftime('%H:%M'),
                               f.partenaire.nom, f.type_client.nom, f.nom_client,
                               f.plaque, f.type_vehicule.nom, f.type_intervention.nom,
                               f.commentaire or '', f.auteur.username])
                for col in ws.columns:
                    ws.column_dimensions[col[0].column_letter].width = max(len(str(c.value or '')) for c in col) + 4
                buf = io.BytesIO()
                wb.save(buf); buf.seek(0)
                return send_file(buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                                 as_attachment=True,
                                 download_name=f'autotrack_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx')
            except ImportError:
                flash('openpyxl non installé. Utilise le format CSV.', 'warning')

    return render_template('export.html', partenaires=partenaires, types_client=types_client)

# ── Admin ─────────────────────────────────────────────────────────────────────
@app.route('/admin')
@login_required
@admin_required
def admin():
    return render_template('admin.html',
        partenaires=Partenaire.query.all(),
        types_client=TypeClient.query.all(),
        types_vehicule=TypeVehicule.query.all(),
        interventions=TypeIntervention.query.all(),
        users=User.query.all()
    )

# Partenaires
@app.route('/admin/partenaire', methods=['POST'])
@login_required
@admin_required
def admin_partenaire():
    action = request.form.get('action')
    if action == 'add':
        nom = request.form['nom'].strip()
        if nom and not Partenaire.query.filter_by(nom=nom).first():
            db.session.add(Partenaire(nom=nom)); db.session.commit()
            flash(f'Partenaire "{nom}" ajouté.', 'success')
    elif action == 'toggle':
        p = Partenaire.query.get(request.form['id'])
        if p: p.actif = not p.actif; db.session.commit()
    elif action == 'delete':
        p = Partenaire.query.get(request.form['id'])
        if p and not p.fiches: db.session.delete(p); db.session.commit(); flash('Partenaire supprimé.', 'success')
        else: flash('Impossible : des fiches existent pour ce partenaire.', 'danger')
    elif action == 'edit':
        p = Partenaire.query.get(request.form['id'])
        if p: p.nom = request.form['nom'].strip(); db.session.commit(); flash('Partenaire modifié.', 'success')
    return redirect(url_for('admin'))

# Types client
@app.route('/admin/type-client', methods=['POST'])
@login_required
@admin_required
def admin_type_client():
    action = request.form.get('action')
    if action == 'add':
        nom = request.form['nom'].strip()
        if nom and not TypeClient.query.filter_by(nom=nom).first():
            db.session.add(TypeClient(nom=nom)); db.session.commit()
            flash(f'Type client "{nom}" ajouté.', 'success')
    elif action == 'toggle':
        t = TypeClient.query.get(request.form['id'])
        if t: t.actif = not t.actif; db.session.commit()
    elif action == 'delete':
        t = TypeClient.query.get(request.form['id'])
        if t and not t.fiches: db.session.delete(t); db.session.commit(); flash('Type client supprimé.', 'success')
        else: flash('Impossible : des fiches existent pour ce type client.', 'danger')
    elif action == 'edit':
        t = TypeClient.query.get(request.form['id'])
        if t: t.nom = request.form['nom'].strip(); db.session.commit(); flash('Type client modifié.', 'success')
    return redirect(url_for('admin'))

# Types véhicule
@app.route('/admin/type-vehicule', methods=['POST'])
@login_required
@admin_required
def admin_type_vehicule():
    action = request.form.get('action')
    if action == 'add':
        nom = request.form['nom'].strip()
        if nom and not TypeVehicule.query.filter_by(nom=nom).first():
            db.session.add(TypeVehicule(nom=nom)); db.session.commit()
            flash(f'Type véhicule "{nom}" ajouté.', 'success')
    elif action == 'toggle':
        t = TypeVehicule.query.get(request.form['id'])
        if t: t.actif = not t.actif; db.session.commit()
    elif action == 'delete':
        t = TypeVehicule.query.get(request.form['id'])
        if t and not t.fiches: db.session.delete(t); db.session.commit(); flash('Type véhicule supprimé.', 'success')
        else: flash('Impossible : des fiches existent pour ce type.', 'danger')
    elif action == 'edit':
        t = TypeVehicule.query.get(request.form['id'])
        if t: t.nom = request.form['nom'].strip(); db.session.commit(); flash('Type véhicule modifié.', 'success')
    return redirect(url_for('admin'))

# Types intervention
@app.route('/admin/type-intervention', methods=['POST'])
@login_required
@admin_required
def admin_type_intervention():
    action = request.form.get('action')
    if action == 'add':
        nom = request.form['nom'].strip()
        tc_id = request.form['type_client_id']
        if nom:
            db.session.add(TypeIntervention(nom=nom, type_client_id=tc_id))
            db.session.commit(); flash(f'Intervention "{nom}" ajoutée.', 'success')
    elif action == 'toggle':
        t = TypeIntervention.query.get(request.form['id'])
        if t: t.actif = not t.actif; db.session.commit()
    elif action == 'delete':
        t = TypeIntervention.query.get(request.form['id'])
        if t and not t.fiches: db.session.delete(t); db.session.commit(); flash('Intervention supprimée.', 'success')
        else: flash('Impossible : des fiches utilisent cette intervention.', 'danger')
    elif action == 'edit':
        t = TypeIntervention.query.get(request.form['id'])
        if t: t.nom = request.form['nom'].strip(); db.session.commit(); flash('Intervention modifiée.', 'success')
    return redirect(url_for('admin'))

# Utilisateurs
@app.route('/admin/user', methods=['POST'])
@login_required
@admin_required
def admin_user():
    action = request.form.get('action')
    if action == 'add':
        username = request.form['username'].strip()
        password = request.form['password']
        is_admin = request.form.get('is_admin') == '1'
        if username and password and not User.query.filter_by(username=username).first():
            u = User(username=username, is_admin=is_admin)
            u.set_password(password)
            db.session.add(u); db.session.commit()
            flash(f'Utilisateur "{username}" créé.', 'success')
        else:
            flash('Nom déjà pris ou champs vides.', 'danger')
    elif action == 'delete':
        u = User.query.get(request.form['id'])
        if u and u.id != current_user.id:
            db.session.delete(u); db.session.commit(); flash('Utilisateur supprimé.', 'success')
    elif action == 'reset_password':
        u = User.query.get(request.form['id'])
        if u: u.set_password(request.form['new_password']); db.session.commit(); flash('Mot de passe réinitialisé.', 'success')
    elif action == 'toggle_admin':
        u = User.query.get(request.form['id'])
        if u and u.id != current_user.id: u.is_admin = not u.is_admin; db.session.commit()
    return redirect(url_for('admin'))

# ── Init DB ──────────────────────────────────────────────────────────────────
def init_db():
    with app.app_context():
        db.create_all()
        if not User.query.filter_by(username='admin').first():
            admin = User(username='admin', is_admin=True)
            admin.set_password('admin123')
            db.session.add(admin)
        if not Partenaire.query.first():
            for n in ['Rauwers', 'Magellan']: db.session.add(Partenaire(nom=n))
        if not TypeClient.query.first():
            for n in ['bpost', 'Autre']: db.session.add(TypeClient(nom=n))
        if not TypeVehicule.query.first():
            for n in ['Voiture', 'Camion', 'Moto', 'Camionnette', 'Autre']:
                db.session.add(TypeVehicule(nom=n))
        db.session.flush()
        if not TypeIntervention.query.first():
            bpost = TypeClient.query.filter_by(nom='bpost').first()
            autre = TypeClient.query.filter_by(nom='Autre').first()
            for n in ['Dépannage', 'Montage', 'Démontage']:
                db.session.add(TypeIntervention(nom=n, type_client_id=bpost.id))
            for n in ['Dépannage', 'Montage', 'Démontage', 'Révision', 'Diagnostic']:
                db.session.add(TypeIntervention(nom=n, type_client_id=autre.id))
        db.session.commit()
        print("✅ Base de données initialisée.")
        print("👤 Admin : admin / admin123")

if __name__ == '__main__':
    init_db()
    port = int(os.environ.get('PORT', 8080))
    app.run(debug=False, host='0.0.0.0', port=port)
